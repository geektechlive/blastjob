"""ATS/Workday-safety tests for exporters/ats.py and exporters/docx.py.

Covers the confirmed defects described in docs/how-it-works/tailoring.md
(V13 / W5): banned characters surviving into the .ats.txt (the exact set behind the
logged Workday rejection, 2026-09-10), ats.py raising instead of being fail-soft like
docx.py/pdf.py, mojibake (U+FFFD) reaching either exported file, and the 2.5MB
Greenhouse parse-limit guard.
"""

import logging

from docx import Document

from blastjob.exporters.ats import _ASCII_MAP, _strip_banned_chars
from blastjob.exporters.ats import write as ats_write
from blastjob.exporters.docx import write as docx_write

_BANNED = '<>[]{}"\\'


def _docx_text(path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


# --- banned characters (D6 / the logged Workday rejection) -----------------------


def test_ats_write_strips_all_banned_characters(tmp_path):
    content = (
        "## Experience\n\n"
        '- Cut latency <100ms> and grew revenue {20%} using [Docker] and "Kubernetes"\\prod\n'
    )
    out = ats_write(content, tmp_path)
    text = out.read_text()
    for ch in _BANNED:
        assert ch not in text, f"banned char {ch!r} survived into .ats.txt"


def test_ats_write_replaces_double_quotes_with_single_quotes(tmp_path):
    out = ats_write('Led the "Alpha" initiative.', tmp_path)
    text = out.read_text()
    assert '"' not in text
    assert "'Alpha'" in text


def test_strip_banned_chars_unit():
    result = _strip_banned_chars("""<a> [b] {c} "d" \\e""")
    assert result == "a b c 'd' e"


def test_ats_write_markdown_link_becomes_bare_url_no_brackets(tmp_path):
    content = "Portfolio: [alexrivera.example.com](https://alexrivera.example.com/work)"
    out = ats_write(content, tmp_path)
    text = out.read_text()
    assert "https://alexrivera.example.com/work" in text
    assert "[" not in text
    assert "]" not in text


# --- fail-soft ASCII contract (D2 / V13) ------------------------------------------


def test_ats_write_never_raises_on_char_outside_ascii_map(tmp_path):
    # CJK, an emoji, and a lone combining mark are not in _ASCII_MAP.
    content = "Summary: \u4e2d\u6587 \U0001f3af experience \u0301 here"
    out = ats_write(content, tmp_path)  # must not raise
    text = out.read_text()
    assert text.isascii()


def test_ats_write_curly_quote_em_dash_unicode_bullet_mix_is_pure_ascii(tmp_path):
    content = (
        "## Experience\n\n"
        "### Director \u2014 Ops\n\n"
        "- Grew \u201crevenue\u201d by 20% \u2022 exceeded every target \u2013 significantly\n"
    )
    out = ats_write(content, tmp_path)
    text = out.read_text()
    assert text.isascii()
    assert "\u2014" not in text
    assert "\u201c" not in text and "\u201d" not in text
    assert "\u2022" not in text


def test_ats_write_u_fffd_never_reaches_output(tmp_path):
    content = "Built the platform \ufffd\ufffd end to end."
    out = ats_write(content, tmp_path)
    text = out.read_text()
    assert "\ufffd" not in text


# --- size guard (2.5MB Greenhouse parse limit) ------------------------------------


def test_ats_warn_if_oversized(tmp_path, caplog):
    from blastjob.exporters.ats import _warn_if_oversized

    big = tmp_path / "big.txt"
    big.write_bytes(b"x" * (2_500_001))
    with caplog.at_level(logging.WARNING):
        _warn_if_oversized(big)
    assert any("parse limit" in r.message for r in caplog.records)


def test_ats_no_warn_under_size_limit(tmp_path, caplog):
    from blastjob.exporters.ats import _warn_if_oversized

    small = tmp_path / "small.txt"
    small.write_bytes(b"x" * 100)
    with caplog.at_level(logging.WARNING):
        _warn_if_oversized(small)
    assert not caplog.records


def test_ats_write_warns_but_does_not_raise_when_over_configured_limit(
    tmp_path, caplog, monkeypatch
):
    import blastjob.exporters.ats as ats_mod

    monkeypatch.setattr(ats_mod, "_MAX_BYTES", 10)
    with caplog.at_level(logging.WARNING):
        dest = ats_mod.write("A resume body well over ten bytes long.", tmp_path)
    assert dest.exists()
    assert any("parse limit" in r.message for r in caplog.records)


def test_docx_warn_if_oversized(tmp_path, caplog):
    from blastjob.exporters.docx import _warn_if_oversized

    big = tmp_path / "big.docx"
    big.write_bytes(b"x" * (2_500_001))
    with caplog.at_level(logging.WARNING):
        _warn_if_oversized(big)
    assert any("parse limit" in r.message for r in caplog.records)


def test_docx_write_warns_but_does_not_raise_when_over_configured_limit(
    tmp_path, caplog, monkeypatch
):
    import blastjob.exporters.docx as docx_mod

    monkeypatch.setattr(docx_mod, "_MAX_BYTES", 10)
    content = "# Name\ncontact@example.com\n\n## Experience\n\n- Did a thing.\n"
    with caplog.at_level(logging.WARNING):
        dest = docx_mod.write(content, tmp_path)
    assert dest.exists()
    assert any("parse limit" in r.message for r in caplog.records)


# --- docx typography normalization + U+FFFD guard (D9) ---------------------------


def test_docx_typography_normalized_like_ats(tmp_path):
    content = (
        "# Alex Rivera\n"
        "alex@example.com\n\n"
        "## Summary\n\n"
        "Led teams \u2014 delivered results \u201cahead of plan\u201d with \u2013 focus.\n"
    )
    out = docx_write(content, tmp_path)
    text = _docx_text(out)
    assert "\u2014" not in text
    assert "\u201c" not in text and "\u201d" not in text
    assert "\u2013" not in text
    # normalized the same way ats.py maps these
    assert _ASCII_MAP["\u2014"] in text


def test_docx_u_fffd_never_reaches_document(tmp_path):
    content = (
        "# Alex Rivera\n"
        "alex@example.com\n\n"
        "## Summary\n\n"
        "Shipped the \ufffd\ufffd migration end to end.\n"
    )
    out = docx_write(content, tmp_path)
    text = _docx_text(out)
    assert "\ufffd" not in text


def test_docx_bullet_formatting_still_uses_bullet_glyph(tmp_path):
    """The renderer's own decorative list bullet is not degraded by the
    typography normalization -- it must still be a real Unicode bullet glyph
    in the visible Word document, not stripped or replaced."""
    content = "# Alex Rivera\ncontact@example.com\n\n## Experience\n\n- Did a thing.\n"
    out = docx_write(content, tmp_path)
    doc = Document(str(out))
    bullet_para = next(p for p in doc.paragraphs if "Did a thing" in p.text)
    assert bullet_para.text.startswith("\u2022")


def test_docx_write_still_returns_correct_path(tmp_path):
    out = docx_write("# Name\ncontact@example.com\n", tmp_path, stem="cover_letter")
    assert out == tmp_path / "cover_letter.docx"
    assert out.exists()
