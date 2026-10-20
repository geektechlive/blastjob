from pathlib import Path

from blastjob.importers import docx as docx_importer


def _table_resume(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jordan Example")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Acme Widgets, Director of IT"
    table.cell(0, 1).text = "2015 - 2020"
    table.cell(1, 0).text = "Skills: Kubernetes, Terraform"
    table.cell(1, 1).text = "Northwind Traders"
    doc.sections[0].header.paragraphs[0].text = "jordan@example.com"
    doc.save(str(path))
    return path


def test_docx_extract_includes_table_cells_and_header(tmp_path):
    path = _table_resume(tmp_path / "resume.docx")

    text = docx_importer.extract(path)

    assert "Jordan Example" in text
    assert "Acme Widgets, Director of IT" in text
    assert "Skills: Kubernetes, Terraform" in text
    assert "jordan@example.com" in text


def test_table_only_docx_is_not_empty(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Only in a table"
    path = tmp_path / "t.docx"
    doc.save(str(path))

    assert docx_importer.extract(path).strip() == "Only in a table"
