"""Read a generated .docx back into text.

Why this exists: the ATS never sees `resume_md`. It receives the `.docx` that
`exporters/docx.py` produces, and that renderer can drop text into a header,
footer, or table without any markdown-level check noticing (the writer never
emits tables today, but a future change could, and a header/footer regression
would be invisible to a pure-markdown check). Structure checks in
`core.resume_format` need to see what actually shipped, not what was asked
for -- "verify the artifact, not the markdown" (packet-quality-program v2,
revision V1).
"""

from __future__ import annotations

from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path


class DocxRegion(StrEnum):
    """Where a line of text sits in the .docx.

    Most parsers read the body; text that ends up in a header, footer, or
    table is a documented parse-failure cause (Textkernel flags header/footer
    contact blocks and table layouts as risk signals) even though it renders
    fine in Word. A resume section that lands in one of those is invisible or
    misread by an ATS.
    """

    BODY = "body"
    HEADER = "header"
    FOOTER = "footer"
    TABLE = "table"


@dataclass(frozen=True)
class DocxLine:
    """One non-empty line of text, tagged with where it lives."""

    text: str
    region: DocxRegion


def extract_lines(path: Path) -> list[DocxLine]:
    """Every non-empty line of text in `path`, tagged with its region.

    Body paragraphs and tables are walked in document order by iterating
    `document.element.body` directly -- python-docx's `document.paragraphs`
    and `document.tables` properties each iterate their own kind only and do
    not preserve the interleaved order, which matters if a future export
    puts a table between two paragraphs. Header/footer paragraphs are read
    from every section, since a multi-section document can define more than
    one of each (this writer only ever produces one section, but a reader
    should not assume that of every .docx it is handed).
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    lines: list[DocxLine] = []

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(DocxLine(text=text, region=DocxRegion.BODY))
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        cell_text = para.text.strip()
                        if cell_text:
                            lines.append(DocxLine(text=cell_text, region=DocxRegion.TABLE))

    for section in doc.sections:
        for para in section.header.paragraphs:
            header_text = para.text.strip()
            if header_text:
                lines.append(DocxLine(text=header_text, region=DocxRegion.HEADER))
        for para in section.footer.paragraphs:
            footer_text = para.text.strip()
            if footer_text:
                lines.append(DocxLine(text=footer_text, region=DocxRegion.FOOTER))

    return lines


def extract_text(path: Path) -> list[str]:
    """Body paragraph text only, in document order.

    This is the read-back form `core.resume_format.check_structure` consumes
    as `docx_lines`. Header/footer/table text is deliberately excluded: a
    body-only reader (most ATS parsers, by Textkernel's own documentation) is
    exactly what this simulates. Use `extract_lines` when you need to know
    THAT some text landed somewhere a body-only reader would never see.
    """
    return [line.text for line in extract_lines(path) if line.region is DocxRegion.BODY]
